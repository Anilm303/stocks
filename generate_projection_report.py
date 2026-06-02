import os
import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# Define Color Palette (Premium Navy & Accents)
NAVY_HEADER_FILL = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
ACCENT_BLUE_FILL = PatternFill(start_color="E8F1F5", end_color="E8F1F5", fill_type="solid")
ZEBRA_FILL = PatternFill(start_color="F5F7FA", end_color="F5F7FA", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

GREEN_TEXT_COLOR = "155724"
GREEN_FILL = PatternFill(start_color="EBF7EE", end_color="EBF7EE", fill_type="solid")

RED_TEXT_COLOR = "721C24"
RED_FILL = PatternFill(start_color="FDEDEC", end_color="FDEDEC", fill_type="solid")

# Fonts
FONT_TITLE = Font(name="Segoe UI", size=16, bold=True, color="1B365D")
FONT_SECTION = Font(name="Segoe UI", size=12, bold=True, color="1B365D")
FONT_HEADER = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
FONT_BODY = Font(name="Segoe UI", size=10, color="000000")
FONT_BODY_BOLD = Font(name="Segoe UI", size=10, bold=True, color="000000")
FONT_KPI_VAL = Font(name="Segoe UI", size=20, bold=True, color="1B365D")
FONT_KPI_LBL = Font(name="Segoe UI", size=9, italic=True, color="555555")

# Borders
BORDER_THIN = Border(
    left=Side(style="thin", color="D3D3D3"),
    right=Side(style="thin", color="D3D3D3"),
    top=Side(style="thin", color="D3D3D3"),
    bottom=Side(style="thin", color="D3D3D3")
)
BORDER_TOP_THIN_BOTTOM_DOUBLE = Border(
    top=Side(style="thin", color="000000"),
    bottom=Side(style="double", color="000000")
)

# Number Formats
FMT_CURRENCY = '"Rs. "#,##0.00'
FMT_INTEGER = "#,##0"
FMT_PERCENT = "0.00%"
FMT_DATE = "YYYY-MM-DD"

def style_sheet(ws, title_text, start_row=4, is_trades_sheet=False):
    """Applies basic premium formatting and column autofit to a sheet."""
    # Ensure grid lines are visible
    ws.views.sheetView[0].showGridLines = True
    
    # Title Block
    ws["A1"] = title_text
    ws["A1"].font = FONT_TITLE
    ws.row_dimensions[1].height = 25
    
    max_row = ws.max_row
    max_col = ws.max_column
    
    # Format Headers (assumed to be at start_row)
    for col_idx in range(1, max_col + 1):
        cell = ws.cell(row=start_row, column=col_idx)
        cell.font = FONT_HEADER
        cell.fill = NAVY_HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[start_row].height = 26

    # Format Data Rows
    for row_idx in range(start_row + 1, max_row + 1):
        is_even = (row_idx % 2 == 0)
        row_fill = ZEBRA_FILL if is_even else WHITE_FILL
        ws.row_dimensions[row_idx].height = 19
        
        # Check if the row contains totals or summary info (e.g. if column A starts with 'Total' or is empty)
        first_cell_val = str(ws.cell(row=row_idx, column=1).value or "")
        is_total_row = "Total" in first_cell_val or first_cell_val == ""
        
        for col_idx in range(1, max_col + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if not is_total_row:
                cell.font = FONT_BODY
                cell.fill = row_fill
                cell.border = BORDER_THIN
            else:
                cell.font = FONT_BODY_BOLD
                cell.border = BORDER_TOP_THIN_BOTTOM_DOUBLE
                
            # Formatting based on content type
            val = cell.value
            header = str(ws.cell(row=start_row, column=col_idx).value).lower()
            
            # Alignments and number formats
            if isinstance(val, (int, float)):
                if "price" in header or "value" in header or "equity" in header or "cash" in header or "total" in header or "pnl" in header or "difference" in header or "allocation" in header:
                    cell.number_format = FMT_CURRENCY
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif "quantity" in header or "trades" in header or "count" in header:
                    cell.number_format = FMT_INTEGER
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                elif "pct" in header or "rate" in header or "return" in header or "cagr" in header:
                    cell.number_format = FMT_PERCENT
                    cell.alignment = Alignment(horizontal="right", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="right", vertical="center")
            elif isinstance(val, str):
                if "-" in val and len(val) == 10 and val.replace("-", "").isdigit(): # Date string
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                elif header == "symbol" or header == "side":
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")

            # Apply soft colors for Buy vs Sell signals if it's the combined trades sheet
            if is_trades_sheet and header == "side":
                if val == "BUY":
                    cell.fill = GREEN_FILL
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color=GREEN_TEXT_COLOR)
                elif val in ["SELL", "FORCED_EXIT"]:
                    cell.fill = RED_FILL
                    cell.font = Font(name="Segoe UI", size=10, bold=True, color=RED_TEXT_COLOR)

    # Auto-fit columns
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Skip title row and empty cells
            if cell.row < start_row or cell.value is None:
                continue
            val_str = str(cell.value)
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

def create_kpi_dashboard(ws, stats):
    """Creates a beautiful KPI card layout on the Summary sheet."""
    ws.views.sheetView[0].showGridLines = False
    
    # Title Block
    ws["A1"] = "Nepali Stocks Trading Backtest & Projection Summary"
    ws["A1"].font = FONT_TITLE
    
    # Section Header
    ws["A3"] = "Performance Dashboard"
    ws["A3"].font = FONT_SECTION
    
    # KPI 1: Initial Cash
    ws["A5"] = "Initial Capital"
    ws["A5"].font = FONT_KPI_LBL
    ws["A6"] = stats["Initial Cash"]
    ws["A6"].font = FONT_KPI_VAL
    ws["A6"].number_format = FMT_CURRENCY
    ws["A5"].fill = ACCENT_BLUE_FILL
    ws["A6"].fill = ACCENT_BLUE_FILL
    
    # KPI 2: Target Portfolio
    ws["C5"] = "Target Portfolio (Projected)"
    ws["C5"].font = FONT_KPI_LBL
    ws["C6"] = stats["Target Capital"]
    ws["C6"].font = FONT_KPI_VAL
    ws["C6"].number_format = FMT_CURRENCY
    ws["C5"].fill = ACCENT_BLUE_FILL
    ws["C6"].fill = ACCENT_BLUE_FILL
    
    # KPI 3: Actual Equity Achieved
    ws["E5"] = "Actual Equity Achieved"
    ws["E5"].font = FONT_KPI_LBL
    ws["E6"] = stats["Actual Equity"]
    ws["E6"].font = FONT_KPI_VAL
    ws["E6"].number_format = FMT_CURRENCY
    ws["E5"].fill = ACCENT_BLUE_FILL
    ws["E6"].fill = ACCENT_BLUE_FILL

    # KPI 4: Total Trades
    ws["A9"] = "Total Trades Run"
    ws["A9"].font = FONT_KPI_LBL
    ws["A10"] = stats["Total Trades"]
    ws["A10"].font = FONT_KPI_VAL
    ws["A10"].number_format = FMT_INTEGER
    ws["A9"].fill = ACCENT_BLUE_FILL
    ws["A10"].fill = ACCENT_BLUE_FILL
    
    # KPI 5: Required CAGR vs Actual
    ws["C9"] = "Required CAGR vs Actual"
    ws["C9"].font = FONT_KPI_LBL
    ws["C10"] = f"{stats['Required CAGR']:.2f}% / {stats['Actual CAGR']:.2f}%"
    ws["C10"].font = FONT_KPI_VAL
    ws["C9"].fill = ACCENT_BLUE_FILL
    ws["C10"].fill = ACCENT_BLUE_FILL

    # KPI 6: Actual Total Return
    ws["E9"] = "Actual Return"
    ws["E9"].font = FONT_KPI_LBL
    ws["E10"] = stats["Actual Return"]
    ws["E10"].font = FONT_KPI_VAL
    ws["E10"].number_format = FMT_PERCENT
    ws["E9"].fill = ACCENT_BLUE_FILL
    ws["E10"].fill = ACCENT_BLUE_FILL

    # KPI Border boxes
    def draw_box(ws, start_c, end_c, start_r, end_r):
        thin_side = Side(border_style="thin", color="B0C4DE")
        for r in range(start_r, end_r + 1):
            for c in range(start_c, end_c + 1):
                cell = ws.cell(row=r, column=c)
                # apply border sides
                t = thin_side if r == start_r else None
                b = thin_side if r == end_r else None
                l = thin_side if c == start_c else None
                rg = thin_side if c == end_c else None
                cell.border = Border(top=t, bottom=b, left=l, right=rg)

    draw_box(ws, 1, 2, 5, 6)
    draw_box(ws, 3, 4, 5, 6)
    draw_box(ws, 5, 6, 5, 6)
    draw_box(ws, 1, 2, 9, 10)
    draw_box(ws, 3, 4, 9, 10)
    draw_box(ws, 5, 6, 9, 10)
    
    # Column width formatting for KPI sheet
    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 15
    ws.column_dimensions["C"].width = 25
    ws.column_dimensions["D"].width = 15
    ws.column_dimensions["E"].width = 22
    ws.column_dimensions["F"].width = 15

def main():
    print("Loading data...")
    # 1. Load Projection Data
    xl_proj = pd.ExcelFile("projection_200k_to_1.5m_new.xlsx")
    df_proj = pd.read_excel(xl_proj, "Projection")
    df_proj_summary = pd.read_excel(xl_proj, "Summary")
    
    # 2. Load Trading Backtest Data
    xl_trade = pd.ExcelFile("trading_report.xlsx")
    df_equity = pd.read_excel(xl_trade, "EquityCurve")
    df_trades = pd.read_excel(xl_trade, "Trades")
    df_buy = pd.read_excel(xl_trade, "BuyTrades")
    df_sell = pd.read_excel(xl_trade, "SellTrades")

    # 3. Process Yearly Actuals from Backtest
    # Extract last record of each year from EquityCurve
    df_equity["Year"] = pd.to_datetime(df_equity["Date"]).dt.year
    df_yearly_equity = df_equity.sort_values("Date").groupby("Year").last().reset_index()
    
    # Calculate yearly buy and sell totals from Trades
    df_trades["Year"] = pd.to_datetime(df_trades["Date"]).dt.year
    yearly_trade_totals = df_trades.groupby(["Year", "Side"])["TradeValue"].sum().unstack().fillna(0).reset_index()
    if "BUY" not in yearly_trade_totals.columns:
        yearly_trade_totals["BUY"] = 0.0
    if "SELL" not in yearly_trade_totals.columns:
        yearly_trade_totals["SELL"] = 0.0
        
    # 4. Merge Projection and Actuals
    # Rename columns to match nicely
    df_proj = df_proj.rename(columns={"Date": "ProjectedDate", "Value": "ProjectedValue"})
    df_yearly_equity = df_yearly_equity.rename(columns={"Date": "ActualDate", "Equity": "ActualEquity", "Cash": "ActualCash"})
    
    merged = pd.merge(df_proj, df_yearly_equity[["Year", "ActualDate", "ActualEquity", "ActualCash"]], on="Year", how="left")
    merged = pd.merge(merged, yearly_trade_totals[["Year", "BUY", "SELL"]], on="Year", how="left").fillna(0)
    
    # Compute PnL and Differences
    merged["YearlyPnL"] = merged["SELL"] - merged["BUY"]
    merged["Difference"] = merged["ActualEquity"] - merged["ProjectedValue"]
    
    # Add Total Row
    total_row = {
        "Year": "Total / Final",
        "ProjectedDate": "",
        "ProjectedValue": df_proj["ProjectedValue"].iloc[-1],
        "ActualDate": "",
        "ActualEquity": df_yearly_equity["ActualEquity"].iloc[-1],
        "ActualCash": df_yearly_equity["ActualCash"].iloc[-1],
        "BUY": merged["BUY"].sum(),
        "SELL": merged["SELL"].sum(),
        "YearlyPnL": merged["YearlyPnL"].sum(),
        "Difference": df_yearly_equity["ActualEquity"].iloc[-1] - df_proj["ProjectedValue"].iloc[-1]
    }
    df_yearly_comp = pd.concat([merged, pd.DataFrame([total_row])], ignore_index=True)

    # 5. Extract KPI summary stats
    initial_cash = float(df_proj_summary["Initial"].iloc[0])
    target_cash = float(df_proj_summary["Target"].iloc[0])
    actual_equity = float(df_yearly_equity["ActualEquity"].iloc[-1])
    required_cagr = float(df_proj_summary["RequiredCAGR_pct"].iloc[0])
    
    years = float(df_proj_summary["Years"].iloc[0])
    actual_cagr = ((actual_equity / initial_cash) ** (1 / years) - 1) * 100
    actual_return = (actual_equity / initial_cash) - 1
    total_trades = len(df_trades)
    
    stats = {
        "Initial Cash": initial_cash,
        "Target Capital": target_cash,
        "Actual Equity": actual_equity,
        "Required CAGR": required_cagr,
        "Actual CAGR": actual_cagr,
        "Actual Return": actual_return,
        "Total Trades": total_trades
    }

    # 6. Generate combined sheet: projection_200k_to_1.5m_buy_sell.xlsx
    print("Writing master report projection_200k_to_1.5m_buy_sell.xlsx...")
    with pd.ExcelWriter("projection_200k_to_1.5m_buy_sell.xlsx", engine="openpyxl") as writer:
        # Create empty sheets first to control order and styling
        pd.DataFrame().to_excel(writer, sheet_name="Summary", index=False)
        df_yearly_comp.to_excel(writer, sheet_name="Yearly_Comparison", index=False, startrow=3)
        df_buy.to_excel(writer, sheet_name="BuyTrades", index=False, startrow=3)
        df_sell.to_excel(writer, sheet_name="SellTrades", index=False, startrow=3)
        df_trades.drop(columns=["Year"], errors="ignore").to_excel(writer, sheet_name="Trades", index=False, startrow=3)
        df_proj.to_excel(writer, sheet_name="Projection", index=False, startrow=3)

    # Format the workbook
    wb = load_workbook("projection_200k_to_1.5m_buy_sell.xlsx")
    create_kpi_dashboard(wb["Summary"], stats)
    style_sheet(wb["Yearly_Comparison"], "Yearly Projection Target vs. Backtest Actuals Comparison", start_row=4)
    style_sheet(wb["BuyTrades"], "Stock Backtest Buy Transactions Detail", start_row=4)
    style_sheet(wb["SellTrades"], "Stock Backtest Sell Transactions Detail", start_row=4)
    style_sheet(wb["Trades"], "All Transactions History Chronological", start_row=4, is_trades_sheet=True)
    style_sheet(wb["Projection"], "Compounded Growth Projection (Rs. 200,000 to Rs. 1,500,000)", start_row=4)
    wb.save("projection_200k_to_1.5m_buy_sell.xlsx")

    # 7. Generate separate Buy file
    print("Writing buy-only report projection_200k_to_1.5m_buy.xlsx...")
    with pd.ExcelWriter("projection_200k_to_1.5m_buy.xlsx", engine="openpyxl") as writer:
        pd.DataFrame().to_excel(writer, sheet_name="Summary", index=False)
        df_yearly_comp.to_excel(writer, sheet_name="Yearly_Comparison", index=False, startrow=3)
        df_buy.to_excel(writer, sheet_name="BuyTrades", index=False, startrow=3)
    
    wb_buy = load_workbook("projection_200k_to_1.5m_buy.xlsx")
    create_kpi_dashboard(wb_buy["Summary"], stats)
    style_sheet(wb_buy["Yearly_Comparison"], "Yearly Projection Target vs. Backtest Actuals Comparison", start_row=4)
    style_sheet(wb_buy["BuyTrades"], "Stock Backtest Buy Transactions Detail", start_row=4)
    wb_buy.save("projection_200k_to_1.5m_buy.xlsx")

    # 8. Generate separate Sell file
    print("Writing sell-only report projection_200k_to_1.5m_sell.xlsx...")
    with pd.ExcelWriter("projection_200k_to_1.5m_sell.xlsx", engine="openpyxl") as writer:
        pd.DataFrame().to_excel(writer, sheet_name="Summary", index=False)
        df_yearly_comp.to_excel(writer, sheet_name="Yearly_Comparison", index=False, startrow=3)
        df_sell.to_excel(writer, sheet_name="SellTrades", index=False, startrow=3)
        
    wb_sell = load_workbook("projection_200k_to_1.5m_sell.xlsx")
    create_kpi_dashboard(wb_sell["Summary"], stats)
    style_sheet(wb_sell["Yearly_Comparison"], "Yearly Projection Target vs. Backtest Actuals Comparison", start_row=4)
    style_sheet(wb_sell["SellTrades"], "Stock Backtest Sell Transactions Detail", start_row=4)
    wb_sell.save("projection_200k_to_1.5m_sell.xlsx")

    print("Successfully generated all reports!")

if __name__ == "__main__":
    main()
