import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

wb = load_workbook("projection_1.5m_buy_sell.xlsx")

# Styles
NAVY_HEADER = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
ZEBRA = PatternFill(start_color="F5F7FA", end_color="F5F7FA", fill_type="solid")
WHITE = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

FONT_HEADER = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
FONT_BODY = Font(name="Segoe UI", size=10, color="000000")

BORDER_THIN = Border(left=Side(style="thin", color="D3D3D3"),
                     right=Side(style="thin", color="D3D3D3"),
                     top=Side(style="thin", color="D3D3D3"),
                     bottom=Side(style="thin", color="D3D3D3"))

def style_ws(ws):
    max_row = ws.max_row
    max_col = ws.max_column
    
    # Header
    for col_idx in range(1, max_col + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = FONT_HEADER
        cell.fill = NAVY_HEADER
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[1].height = 26
    
    # Body
    for r in range(2, max_row + 1):
        fill = ZEBRA if r % 2 == 0 else WHITE
        ws.row_dimensions[r].height = 19
        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = FONT_BODY
            cell.fill = fill
            cell.border = BORDER_THIN
            header = str(ws.cell(row=1, column=c).value).lower()
            if isinstance(cell.value, (int, float)):
                if any(x in header for x in ["price", "value", "pnl", "cash", "allocated", "allocation"]):
                    cell.number_format = '"Rs." #,##0.00'
                else:
                    cell.number_format = '#,##0'

    # Auto-fit
    for col in ws.columns:
        col_letter = col[0].column_letter
        max_len = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

for sheetname in wb.sheetnames:
    style_ws(wb[sheetname])

wb.save("projection_1.5m_buy_sell.xlsx")
print("Styling complete!")
