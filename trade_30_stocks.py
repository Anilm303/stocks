from __future__ import annotations

import argparse
from dataclasses import dataclass
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


INITIAL_CASH_DEFAULT = 200_000.0
DEFAULT_SMA_WINDOW = 20
DEFAULT_RSI_BUY_MIN = 45.0
DEFAULT_RSI_BUY_MAX = 70.0
DEFAULT_RSI_SELL_MIN = 35.0
DEFAULT_RSI_SELL_MAX = 75.0


@dataclass
class Position:
    symbol: str
    quantity: int
    entry_price: float
    entry_date: pd.Timestamp


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Backtest a simple trading strategy across up to 30 Nepal stock symbols."
    )
    parser.add_argument("input", help="Path to a CSV file, Excel file, or folder with historical prices.")
    parser.add_argument(
        "--output",
        default="trading_report.xlsx",
        help="Path to the Excel report to create.",
    )
    parser.add_argument(
        "--sheet",
        default=None,
        help="Optional sheet name if the workbook has multiple sheets.",
    )
    parser.add_argument(
        "--initial-cash",
        type=float,
        default=INITIAL_CASH_DEFAULT,
        help="Starting capital in rupees.",
    )
    parser.add_argument(
        "--sma-window",
        type=int,
        default=DEFAULT_SMA_WINDOW,
        help="Simple moving average window used when no Signal column is provided.",
    )
    parser.add_argument(
        "--max-positions",
        type=int,
        default=30,
        help="Maximum number of open positions to hold at once.",
    )
    parser.add_argument(
        "--export-trades",
        action="store_true",
        help="Also export the trades dataframe to trades.csv for analysis.",
    )
    return parser.parse_args()


def load_price_data(path: Path, sheet_name: str | None = None) -> pd.DataFrame:
    if path.is_dir():
        frames = load_from_directory(path)
    elif path.suffix.lower() == ".csv":
        frames = [standardize_frame(pd.read_csv(path), path.stem)]
    else:
        workbook = pd.ExcelFile(path)
        if sheet_name is not None:
            raw_frames = [pd.read_excel(workbook, sheet_name=sheet_name)]
            sheet_names = [sheet_name]
        else:
            sheet_names = workbook.sheet_names
            raw_frames = [pd.read_excel(workbook, sheet_name=name) for name in sheet_names]

        frames = [standardize_frame(frame, current_sheet_name or "Sheet1") for current_sheet_name, frame in zip(sheet_names, raw_frames)]

    data = pd.concat(frames, ignore_index=True)
    data["Date"] = pd.to_datetime(data["Date"], format='mixed')
    data = data.sort_values(["Date", "Symbol"]).reset_index(drop=True)
    return data


def load_from_directory(directory: Path) -> list[pd.DataFrame]:
    frames: list[pd.DataFrame] = []
    for file_path in sorted(directory.iterdir()):
        if file_path.suffix.lower() == ".csv":
            frames.append(standardize_frame(pd.read_csv(file_path), file_path.stem))
        elif file_path.suffix.lower() in {".xlsx", ".xlsm", ".xls"}:
            workbook = pd.ExcelFile(file_path)
            sheet_frames = [pd.read_excel(workbook, sheet_name=name) for name in workbook.sheet_names]
            frames.extend(
                standardize_frame(frame, f"{file_path.stem}_{sheet_name or 'Sheet1'}")
                for sheet_name, frame in zip(workbook.sheet_names, sheet_frames)
            )

    if not frames:
        raise ValueError(f"No CSV or Excel files found in {directory}")
    return frames


def standardize_frame(frame: pd.DataFrame, fallback_symbol: str) -> pd.DataFrame:
    columns = {column.lower().strip(): column for column in frame.columns}

    if "date" not in columns:
        raise ValueError("Each sheet must include a Date column.")

    if "symbol" in columns:
        symbol_series = frame[columns["symbol"]].astype(str)
    else:
        symbol_series = pd.Series([fallback_symbol] * len(frame))

    close_column = columns.get("close") or columns.get("ltp") or columns.get("adj close") or columns.get("price")
    if close_column is None:
        raise ValueError("Each sheet must include a Close, Ltp, Adj Close, or Price column.")

    open_column = columns.get("open")
    signal_column = columns.get("signal")

    result = frame.copy()
    result = result.rename(columns={columns["date"]: "Date"})

    if "symbol" in columns:
        result = result.rename(columns={columns["symbol"]: "Symbol"})
    else:
        result["Symbol"] = symbol_series

    if open_column is not None:
        result = result.rename(columns={open_column: "Open"})
    else:
        result["Open"] = pd.NA

    result = result.rename(columns={close_column: "Close"})

    if signal_column is not None:
        result = result.rename(columns={signal_column: "Signal"})

    result["Open"] = pd.to_numeric(result["Open"], errors="coerce")
    result["Close"] = pd.to_numeric(result["Close"], errors="coerce")
    numeric_columns = [
        "High",
        "Low",
        "Volume",
        "Turnover",
        "Daily_Return",
        "Log_Return",
        "SMA_5",
        "SMA_20",
        "EMA_12",
        "EMA_26",
        "RSI_14",
        "MACD",
        "MACD_Signal",
        "ATR_14",
        "BB_Middle",
        "BB_Std",
        "BB_Upper",
        "BB_Lower",
        "OBV",
    ]
    for column in numeric_columns:
        if column in result.columns:
            result[column] = pd.to_numeric(result[column], errors="coerce")

    result = result.dropna(subset=["Date", "Symbol", "Close"]).copy()
    return result


def validate_input_columns(frame: pd.DataFrame) -> None:
    required_columns = {"date", "symbol", "close"}
    present_columns = {column.lower().strip() for column in frame.columns}
    missing_columns = required_columns - present_columns
    if missing_columns:
        raise ValueError(
            "Your sheet should have these columns: Date, Symbol, Close. Optional columns: Open, Signal. Missing: "
            + ", ".join(sorted(missing_columns))
        )


def add_strategy_signals(data: pd.DataFrame, sma_window: int) -> pd.DataFrame:
    if "Signal" in data.columns:
        signal = data["Signal"]
        if signal.dtype.kind in {"i", "u", "f"}:
            data["StrategySignal"] = signal.fillna(0).astype(int).clip(-1, 1)
        else:
            normalized = signal.astype(str).str.strip().str.lower()
            data["StrategySignal"] = normalized.map(
                {
                    "buy": 1,
                    "long": 1,
                    "1": 1,
                    "sell": -1,
                    "exit": -1,
                    "-1": -1,
                    "hold": 0,
                    "0": 0,
                }
            ).fillna(0).astype(int)
        return data

    indicator_columns = {column.lower().strip() for column in data.columns}
    has_indicator_set = {"sma_20", "rsi_14", "macd", "macd_signal"}.issubset(indicator_columns)

    frames: list[pd.DataFrame] = []
    for symbol, group in data.groupby("Symbol", sort=False):
        ordered = group.sort_values("Date").copy()

        if has_indicator_set:
            ordered["SMA20"] = pd.to_numeric(ordered.get("SMA_20"), errors="coerce")
            ordered["RSI14"] = pd.to_numeric(ordered.get("RSI_14"), errors="coerce")
            ordered["MACD"] = pd.to_numeric(ordered.get("MACD"), errors="coerce")
            ordered["MACDSignal"] = pd.to_numeric(ordered.get("MACD_Signal"), errors="coerce")

            prev_close = ordered["Close"].shift(1)
            prev_sma = ordered["SMA20"].shift(1)
            prev_macd = ordered["MACD"].shift(1)
            prev_macd_signal = ordered["MACDSignal"].shift(1)

            bullish_price_cross = (ordered["Close"] > ordered["SMA20"]) & (prev_close <= prev_sma)
            bullish_macd_cross = (ordered["MACD"] > ordered["MACDSignal"]) & (prev_macd <= prev_macd_signal)
            buy_signal = bullish_price_cross & bullish_macd_cross & ordered["RSI14"].between(DEFAULT_RSI_BUY_MIN, DEFAULT_RSI_BUY_MAX)

            bearish_price_cross = (ordered["Close"] < ordered["SMA20"]) & (prev_close >= prev_sma)
            bearish_macd_cross = (ordered["MACD"] < ordered["MACDSignal"]) & (prev_macd >= prev_macd_signal)
            sell_signal = bearish_price_cross | bearish_macd_cross | (ordered["RSI14"] >= DEFAULT_RSI_SELL_MAX) | (ordered["RSI14"] <= DEFAULT_RSI_SELL_MIN)
        else:
            ordered["SMA"] = ordered["Close"].rolling(window=sma_window, min_periods=sma_window).mean()
            prev_close = ordered["Close"].shift(1)
            prev_sma = ordered["SMA"].shift(1)
            buy_signal = (ordered["Close"] > ordered["SMA"]) & (prev_close <= prev_sma)
            sell_signal = (ordered["Close"] < ordered["SMA"]) & (prev_close >= prev_sma)

        ordered["StrategySignal"] = 0
        ordered.loc[buy_signal, "StrategySignal"] = 1
        ordered.loc[sell_signal, "StrategySignal"] = -1
        frames.append(ordered)

    return pd.concat(frames, ignore_index=True)


def backtest(data: pd.DataFrame, initial_cash: float, max_positions: int) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    cash = float(initial_cash)
    positions: dict[str, Position] = {}
    trade_rows: list[dict[str, object]] = []
    equity_rows: list[dict[str, object]] = []

    for date, day in data.groupby("Date", sort=True):
        day = day.sort_values("Symbol")
        day_prices = {row.Symbol: float(row.Close) for row in day.itertuples(index=False)}

        for row in day.itertuples(index=False):
            if int(row.StrategySignal) != -1:
                continue
            symbol = str(row.Symbol)
            if symbol not in positions:
                continue
            exit_price = float(row.Open) if pd.notna(row.Open) else float(row.Close)
            position = positions.pop(symbol)
            cash += position.quantity * exit_price
            trade_rows.append(
                {
                    "Date": date,
                    "Symbol": symbol,
                    "Side": "SELL",
                    "Quantity": position.quantity,
                    "Price": exit_price,
                    "TradeValue": position.quantity * exit_price,
                    "CashAfterTrade": cash,
                    "EntryDate": position.entry_date,
                    "EntryPrice": position.entry_price,
                }
            )

        buy_candidates = [row for row in day.itertuples(index=False) if int(row.StrategySignal) == 1 and str(row.Symbol) not in positions]
        if buy_candidates:
            slots_available = max(max_positions - len(positions), 0)
            selected_candidates = buy_candidates[:slots_available]

            # Prefer equal per-position allocation based on initial cash
            per_position_budget = float(initial_cash) / float(max_positions) if max_positions > 0 else float(initial_cash)
            m = len(selected_candidates)
            if m == 0:
                allocations: list[float] = []
            else:
                total_needed = per_position_budget * m
                if cash >= total_needed:
                    allocations = [per_position_budget] * m
                else:
                    # not enough cash to fully fund per_position_budget for all; split remaining cash equally
                    allocations = [cash / m] * m

            for row, alloc in zip(selected_candidates, allocations):
                symbol = str(row.Symbol)
                buy_price = float(row.Open) if pd.notna(row.Open) else float(row.Close)
                if buy_price <= 0 or alloc <= 0:
                    continue
                quantity = int(alloc // buy_price)
                if quantity <= 0:
                    continue
                trade_value = quantity * buy_price
                cash -= trade_value
                positions[symbol] = Position(
                    symbol=symbol,
                    quantity=quantity,
                    entry_price=buy_price,
                    entry_date=pd.Timestamp(date),
                )
                trade_rows.append(
                    {
                        "Date": date,
                        "Symbol": symbol,
                        "Side": "BUY",
                        "Quantity": quantity,
                        "Price": buy_price,
                        "TradeValue": trade_value,
                        "Allocated": alloc,
                        "RemainingAllocation": alloc - trade_value,
                        "CashAfterTrade": cash,
                        "EntryDate": pd.Timestamp(date),
                        "EntryPrice": buy_price,
                    }
                )

        unrealized_value = sum(position.quantity * day_prices.get(symbol, position.entry_price) for symbol, position in positions.items())
        equity_rows.append(
            {
                "Date": date,
                "Cash": cash,
                "PositionsValue": unrealized_value,
                "Equity": cash + unrealized_value,
                "OpenPositions": len(positions),
            }
        )

    final_date = data["Date"].max()
    final_prices = (
        data.sort_values("Date")
        .groupby("Symbol", as_index=False)
        .tail(1)
        .set_index("Symbol")["Close"]
        .to_dict()
    )
    for symbol, position in list(positions.items()):
        exit_price = float(final_prices.get(symbol, position.entry_price))
        cash += position.quantity * exit_price
        trade_rows.append(
            {
                "Date": final_date,
                "Symbol": symbol,
                "Side": "FORCED_EXIT",
                "Quantity": position.quantity,
                "Price": exit_price,
                "TradeValue": position.quantity * exit_price,
                "CashAfterTrade": cash,
                "EntryDate": position.entry_date,
                "EntryPrice": position.entry_price,
            }
        )
        positions.pop(symbol, None)

    trade_columns = ["Date", "Symbol", "Side", "Quantity", "Price", "TradeValue", "Allocated", "RemainingAllocation", "CashAfterTrade", "EntryDate", "EntryPrice"]
    trade_df = pd.DataFrame(trade_rows, columns=trade_columns)
    if not trade_df.empty:
        trade_df = trade_df.sort_values(["Date", "Symbol", "Side"]).reset_index(drop=True)
        # Format Date and EntryDate to date-only strings (remove trailing 00:00:00 time)
        for col in ("Date", "EntryDate"):
            if col in trade_df.columns:
                trade_df[col] = pd.to_datetime(trade_df[col], format='mixed').dt.strftime("%Y-%m-%d")
        # Drop exact duplicate trade rows (same Date, Symbol, Side, Quantity, Price, TradeValue)
        dup_subset = [c for c in ("Date", "Symbol", "Side", "Quantity", "Price", "TradeValue") if c in trade_df.columns]
        if dup_subset:
            trade_df = trade_df.drop_duplicates(subset=dup_subset, keep="first").reset_index(drop=True)
    equity_df = pd.DataFrame(equity_rows)
    if not equity_df.empty:
        equity_df = equity_df.sort_values("Date").reset_index(drop=True)
        # Format equity Date to date-only string
        if "Date" in equity_df.columns:
            equity_df["Date"] = pd.to_datetime(equity_df["Date"], format='mixed').dt.strftime("%Y-%m-%d")
    holdings_df = pd.DataFrame(
        [
            {
                "Symbol": position.symbol,
                "Quantity": position.quantity,
                "EntryDate": position.entry_date,
                "EntryPrice": position.entry_price,
            }
            for position in positions.values()
        ]
    )
    # Format holdings EntryDate to date-only string
    if not holdings_df.empty and "EntryDate" in holdings_df.columns:
        holdings_df["EntryDate"] = pd.to_datetime(holdings_df["EntryDate"], format='mixed').dt.strftime("%Y-%m-%d")
    if not positions:
        equity_df = pd.concat(
            [
                equity_df,
                pd.DataFrame(
                    [
                        {
                            "Date": final_date,
                            "Cash": cash,
                            "PositionsValue": 0.0,
                            "Equity": cash,
                            "OpenPositions": 0,
                        }
                    ]
                ),
            ],
            ignore_index=True,
        )

    return trade_df, equity_df, holdings_df


def compute_summary(equity_df: pd.DataFrame, trade_df: pd.DataFrame, initial_cash: float, start_date: pd.Timestamp, end_date: pd.Timestamp) -> pd.DataFrame:
    if equity_df.empty:
        final_equity = initial_cash
        max_drawdown = 0.0
    else:
        final_equity = float(equity_df["Equity"].iloc[-1])
        running_max = equity_df["Equity"].cummax()
        drawdown = equity_df["Equity"] / running_max - 1.0
        max_drawdown = float(drawdown.min())

    years = max((end_date - start_date).days / 365.25, 1 / 365.25)
    cagr = (final_equity / initial_cash) ** (1 / years) - 1 if initial_cash > 0 else 0.0
    total_return = final_equity / initial_cash - 1 if initial_cash > 0 else 0.0

    buys = trade_df[trade_df["Side"] == "BUY"] if not trade_df.empty else pd.DataFrame()
    sells = trade_df[trade_df["Side"].isin(["SELL", "FORCED_EXIT"])] if not trade_df.empty else pd.DataFrame()
    realized_trades = min(len(buys), len(sells))
    win_rate = 0.0
    if realized_trades > 0:
        pnl_by_symbol: dict[str, float] = {}
        for _, trade in trade_df.iterrows():
            symbol = str(trade["Symbol"])
            pnl_by_symbol.setdefault(symbol, 0.0)
            if trade["Side"] == "BUY":
                pnl_by_symbol[symbol] -= float(trade["TradeValue"])
            else:
                pnl_by_symbol[symbol] += float(trade["TradeValue"])
        positive = sum(1 for value in pnl_by_symbol.values() if value > 0)
        win_rate = positive / max(len(pnl_by_symbol), 1)

    return pd.DataFrame(
        [
            {"Metric": "Initial Cash", "Value": initial_cash},
            {"Metric": "Final Equity", "Value": final_equity},
            {"Metric": "Total Return", "Value": total_return},
            {"Metric": "Total Trades", "Value": len(trade_df)},
        ]
    )


def write_workbook(output_path: Path, sheet_frames: dict[str, pd.DataFrame]) -> None:
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        for sheet_name, frame in sheet_frames.items():
            frame.to_excel(writer, index=False, sheet_name=sheet_name)

    workbook = load_workbook(output_path)
    decimal_columns = {"Allocated", "RemainingAllocation"}
    number_format = "0.00"

    for worksheet in workbook.worksheets:
        header_map = {cell.value: cell.column for cell in worksheet[1] if cell.value in decimal_columns}
        for column_index in header_map.values():
            column_letter = get_column_letter(column_index)
            for row_index in range(2, worksheet.max_row + 1):
                worksheet[f"{column_letter}{row_index}"].number_format = number_format

    workbook.save(output_path)


def build_buy_trades_df(trade_df: pd.DataFrame) -> pd.DataFrame:
    if trade_df.empty:
        return trade_df.copy()

    buy_df = trade_df[trade_df["Side"] == "BUY"].copy()
    if buy_df.empty:
        return buy_df

    return buy_df.rename(
        columns={
            "Price": "BuyPrice",
            "TradeValue": "BuyValue",
        }
    )[["Date", "Symbol", "Quantity", "BuyPrice", "BuyValue", "Allocated", "RemainingAllocation", "CashAfterTrade", "EntryDate"]]


def build_sell_trades_df(trade_df: pd.DataFrame) -> pd.DataFrame:
    if trade_df.empty:
        return trade_df.copy()

    sell_df = trade_df[trade_df["Side"].isin(["SELL", "FORCED_EXIT"])].copy()
    if sell_df.empty:
        return sell_df

    sell_df = sell_df.rename(
        columns={
            "Price": "SellPrice",
            "TradeValue": "SellValue",
        }
    )
    sell_df["PnL"] = sell_df["SellValue"] - (sell_df["Quantity"] * sell_df["EntryPrice"])
    return sell_df[["Date", "Symbol", "Quantity", "SellPrice", "SellValue", "PnL", "CashAfterTrade", "EntryDate", "EntryPrice"]]


def write_report(output_path: Path, summary_df: pd.DataFrame, trade_df: pd.DataFrame, equity_df: pd.DataFrame, holdings_df: pd.DataFrame, data_df: pd.DataFrame) -> None:
    buy_df = build_buy_trades_df(trade_df)
    sell_df = build_sell_trades_df(trade_df)

    # Ensure Date and EntryDate columns are formatted as date-only strings across all sheets
    def _format_dates(df: pd.DataFrame) -> pd.DataFrame:
        if df is None or df.empty:
            return df
        df = df.copy()
        for col in ("Date", "EntryDate"):
            if col in df.columns:
                df[col] = pd.to_datetime(df[col], format='mixed').dt.strftime("%Y-%m-%d")
        return df

    buy_df = _format_dates(buy_df)
    sell_df = _format_dates(sell_df)
    trade_df = _format_dates(trade_df)
    equity_df = _format_dates(equity_df)
    holdings_df = _format_dates(holdings_df)
    data_df = _format_dates(data_df)

    combined_sheets = {
        "Summary": summary_df,
        "BuyTrades": buy_df,
        "SellTrades": sell_df,
        "Trades": trade_df,
        "EquityCurve": equity_df,
        "OpenPositions": holdings_df,
        "Signals": data_df,
    }
    write_workbook(output_path, combined_sheets)

    buy_report_path = output_path.with_name(f"{output_path.stem}_buy{output_path.suffix}")
    sell_report_path = output_path.with_name(f"{output_path.stem}_sell{output_path.suffix}")
    buy_sell_report_path = output_path.with_name(f"{output_path.stem}_buy_sell{output_path.suffix}")
    combined_trades_df = pd.concat(
        [
            buy_df.assign(Side="BUY"),
            sell_df.assign(Side="SELL"),
        ],
        ignore_index=True,
        sort=False,
    )

    write_workbook(buy_report_path, {"BuyTrades": buy_df})
    write_workbook(sell_report_path, {"SellTrades": sell_df})
    write_workbook(buy_sell_report_path, {"Trades": combined_trades_df})


def main() -> None:
    args = parse_args()
    input_path = Path(args.input)
    output_path = Path(args.output)

    data = load_price_data(input_path, args.sheet)
    validate_input_columns(data.rename(columns={"Date": "date", "Symbol": "symbol", "Open": "open", "Close": "close", "Signal": "signal"}))
    data = add_strategy_signals(data, args.sma_window)
    trade_df, equity_df, holdings_df = backtest(data, args.initial_cash, args.max_positions)

    start_date = data["Date"].min()
    end_date = data["Date"].max()
    summary_df = compute_summary(equity_df, trade_df, args.initial_cash, start_date, end_date)

    write_report(output_path, summary_df, trade_df, equity_df, holdings_df, data)
    if args.export_trades:
        trade_df.to_csv("trades.csv", index=False)
    print(f"Combined report written to {output_path.resolve()}")
    print(f"Buy report written to {output_path.with_name(f'{output_path.stem}_buy{output_path.suffix}').resolve()}")
    print(f"Sell report written to {output_path.with_name(f'{output_path.stem}_sell{output_path.suffix}').resolve()}")
    print(summary_df.to_string(index=False))


if __name__ == "__main__":
    main()
     